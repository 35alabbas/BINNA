"""
BINNA — Financial Data API  (Phase 6)
Routes: /api/v1/projects/<project_id>/financial
Stores the entire financial data as a JSON blob per project.
"""
import csv
import io
import json
import re
import zipfile
from itertools import islice
from pathlib import Path
from datetime import datetime, timezone

from flask import Blueprint, jsonify, request
from flask_jwt_extended import jwt_required, get_jwt_identity

from app import db
from app.models import FinancialData, Project, ProjectMember, _ROLE_ORDER
from app.tenancy import authorize_project_access
from app.payment_expenses import (
    preserve_linked_expenses,
    preserve_required_expense_categories,
)

financial_bp = Blueprint('financial', __name__)

_BOQ_IMPORT_MAX_BYTES = 5 * 1024 * 1024
_BOQ_IMPORT_MAX_ROWS = 5000
_ARABIC_DIGITS = str.maketrans('٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹', '01234567890123456789')


# This blueprint's per-module permission key (see ProjectMember.can / effective_permissions).
_MODULE = 'financial'


def _check_access(project_id, user_id):
    # Company-scoped gate (READ): returns the project only when it belongs to the
    # caller's company AND the user is a member; otherwise None.
    proj, _mem, allowed = authorize_project_access(
        project_id, user_id, _MODULE, 'view', 'viewer')
    return proj if allowed else None


def _requires_role(project_id, user_id, min_role='member'):
    """Company-scoped gate (WRITE) — mirrors the seven already-gated modules.
    (None, None) → 404 (no access / wrong tenant); (project, None) → 403 (role too low)."""
    proj, mem, allowed = authorize_project_access(
        project_id, user_id, _MODULE, 'edit', min_role)
    if not proj:
        return None, None
    if not allowed:
        return proj, None
    if not allowed:
        return proj, None   # module edit revoked via the per-member permission map → 403
    return proj, mem


@financial_bp.route('', methods=['GET'])
@jwt_required()
def get_financial(project_id):
    uid = int(get_jwt_identity())
    proj = _check_access(project_id, uid)
    if not proj:
        return jsonify({'error': 'Project not found or access denied'}), 404

    fd = FinancialData.query.filter_by(project_id=project_id).first()
    if not fd or not fd.data:
        return jsonify({'data': None}), 200

    try:
        data = json.loads(fd.data)
    except (ValueError, TypeError):
        data = None

    return jsonify({'data': data}), 200


@financial_bp.route('', methods=['PUT'])
@jwt_required()
def save_financial(project_id):
    uid = int(get_jwt_identity())
    proj, mem = _requires_role(project_id, uid, 'member')
    if proj is None:
        return jsonify({'error': 'Project not found'}), 404
    if mem is None:
        return jsonify({'error': 'Insufficient permissions — member role required'}), 403

    body = request.get_json(silent=True) or {}
    payload = body.get('data')

    fd = FinancialData.query.filter_by(project_id=project_id).first()
    if not fd:
        fd = FinancialData(project_id=project_id, updated_by=uid)
        db.session.add(fd)

    existing = {}
    if fd.data:
        try:
            existing = json.loads(fd.data) or {}
        except (ValueError, TypeError):
            existing = {}
    if isinstance(payload, dict):
        payload['expenses'] = preserve_linked_expenses(
            existing.get('expenses', []), payload.get('expenses', []))
        payload['expenseCategories'] = preserve_required_expense_categories(
            existing.get('expenseCategories', []), payload.get('expenseCategories', []))
    fd.company_id = proj.company_id
    fd.data       = json.dumps(payload, ensure_ascii=False)
    fd.updated_by = uid
    fd.updated_at = datetime.now(timezone.utc)
    db.session.commit()
    return jsonify({'ok': True}), 200


@financial_bp.route('/income', methods=['PUT'])
@jwt_required()
def save_income_expenses(project_id):
    """Partial update — only writes income/expense fields, never touches BOQ/VO/payments."""
    uid = int(get_jwt_identity())
    proj, mem = _requires_role(project_id, uid, 'member')
    if proj is None:
        return jsonify({'error': 'Project not found'}), 404
    if mem is None:
        return jsonify({'error': 'Insufficient permissions — member role required'}), 403

    body = request.get_json(silent=True) or {}

    fd = FinancialData.query.filter_by(project_id=project_id).first()
    if not fd:
        fd = FinancialData(project_id=project_id, updated_by=uid)
        db.session.add(fd)

    # Load existing blob, preserving contract/BOQ/VO/payments
    existing = {}
    if fd.data:
        try:
            existing = json.loads(fd.data) or {}
        except (ValueError, TypeError):
            existing = {}

    # Patch only the income/expense fields
    existing['income']             = body.get('income',             existing.get('income', []))
    existing['expenses']           = preserve_linked_expenses(
        existing.get('expenses', []), body.get('expenses', existing.get('expenses', [])))
    existing['incomeCategories']   = body.get('incomeCategories',   existing.get('incomeCategories', []))
    existing['expenseCategories']  = preserve_required_expense_categories(
        existing.get('expenseCategories', []),
        body.get('expenseCategories', existing.get('expenseCategories', [])))
    existing['usdRate']            = body.get('usdRate',            existing.get('usdRate', 1310))

    fd.data       = json.dumps(existing, ensure_ascii=False)
    fd.updated_by = uid
    fd.updated_at = datetime.now(timezone.utc)
    db.session.commit()
    return jsonify({'ok': True}), 200


def _category_lists(payload):
    return {
        'income_categories': payload.get('incomeCategories', []),
        'expense_categories': payload.get('expenseCategories', []),
    }


@financial_bp.route('/categories', methods=['GET'])
@jwt_required()
def get_categories(project_id):
    uid = int(get_jwt_identity())
    proj = _check_access(project_id, uid)
    if not proj:
        return jsonify({'error': 'Project not found or access denied'}), 404
    fd = FinancialData.query.filter_by(project_id=project_id).first()
    try:
        payload = json.loads(fd.data) if fd and fd.data else {}
    except (ValueError, TypeError):
        payload = {}
    return jsonify(_category_lists(payload if isinstance(payload, dict) else {})), 200


@financial_bp.route('/categories', methods=['POST'])
@jwt_required()
def create_category(project_id):
    uid = int(get_jwt_identity())
    proj, mem = _requires_role(project_id, uid, 'member')
    if not proj:
        return jsonify({'error': 'Project not found'}), 404
    if mem is None:
        return jsonify({'error': 'Insufficient permissions'}), 403
    body = request.get_json(silent=True) or {}
    kind = body.get('kind')
    name = (body.get('name') or '').strip()
    if kind not in ('income', 'expense') or not name:
        return jsonify({'error': 'kind and name are required'}), 400
    import uuid
    fd = FinancialData.query.filter_by(project_id=project_id).first()
    if not fd:
        fd = FinancialData(company_id=proj.company_id, project_id=project_id,
                           updated_by=uid, data='{}')
        db.session.add(fd)
    try:
        payload = json.loads(fd.data) if fd.data else {}
    except (ValueError, TypeError):
        payload = {}
    key = 'incomeCategories' if kind == 'income' else 'expenseCategories'
    categories = payload.setdefault(key, [])
    if any(isinstance(category, dict) and
           str(category.get('name') or '').strip().casefold() == name.casefold()
           for category in categories):
        return jsonify({'error': 'Category name already exists'}), 409
    category = {'id': str(uuid.uuid4()), 'name': name, 'archived': False}
    categories.append(category)
    fd.data = json.dumps(payload, ensure_ascii=False)
    db.session.commit()
    return jsonify({'category': category}), 201


@financial_bp.route('/categories/<kind>/<category_id>', methods=['PUT'])
@jwt_required()
def update_category(project_id, kind, category_id):
    uid = int(get_jwt_identity())
    proj, mem = _requires_role(project_id, uid, 'member')
    if not proj:
        return jsonify({'error': 'Project not found'}), 404
    if mem is None:
        return jsonify({'error': 'Insufficient permissions'}), 403
    if kind not in ('income', 'expense'):
        return jsonify({'error': 'Invalid category kind'}), 400
    fd = FinancialData.query.filter_by(project_id=project_id).first()
    if not fd or not fd.data:
        return jsonify({'error': 'Category not found'}), 404
    try:
        payload = json.loads(fd.data) or {}
    except (ValueError, TypeError):
        payload = {}
    key = 'incomeCategories' if kind == 'income' else 'expenseCategories'
    category = next((c for c in payload.get(key, []) if isinstance(c, dict) and str(c.get('id')) == category_id), None)
    if not category:
        return jsonify({'error': 'Category not found'}), 404
    body = request.get_json(silent=True) or {}
    if 'name' in body:
        name = (body.get('name') or '').strip()
        if not name:
            return jsonify({'error': 'name is required'}), 400
        if any(isinstance(item, dict) and str(item.get('id')) != category_id and
               str(item.get('name') or '').strip().casefold() == name.casefold()
               for item in payload.get(key, [])):
            return jsonify({'error': 'Category name already exists'}), 409
        category['name'] = name
    if 'archived' in body:
        if category.get('system_required') and body.get('archived'):
            return jsonify({'error': 'Required payment category cannot be archived'}), 409
        category['archived'] = bool(body.get('archived'))
    fd.data = json.dumps(payload, ensure_ascii=False)
    db.session.commit()
    return jsonify({'category': category}), 200


def _parse_localized_number(value):
    if value is None or value == '':
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text_value = str(value).strip().translate(_ARABIC_DIGITS)
    if not text_value or text_value.startswith('='):
        return None
    text_value = text_value.replace('\u066c', '').replace(',', '').replace(' ', '')
    text_value = text_value.replace('\u066b', '.')
    if not re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)', text_value):
        return None
    try:
        return float(text_value)
    except ValueError:
        return None


def _read_csv_rows(content):
    decoded = None
    for encoding in ('utf-8-sig', 'utf-8', 'cp1256'):
        try:
            decoded = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise ValueError('CSV encoding is not supported; use UTF-8')
    return list(csv.reader(io.StringIO(decoded)))


def _read_xlsx_rows(content):
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    sheet = workbook.active
    rows = [list(row) for row in islice(
        sheet.iter_rows(values_only=True), _BOQ_IMPORT_MAX_ROWS + 2,
    )]
    workbook.close()
    return rows


def _clean_sheet(rows):
    cleaned = []
    for row in rows:
        values = list(row)
        while values and (values[-1] is None or str(values[-1]).strip() == ''):
            values.pop()
        if values:
            cleaned.append(values)
    if not cleaned:
        raise ValueError('The selected file contains no rows')
    if len(cleaned) - 1 > _BOQ_IMPORT_MAX_ROWS:
        raise ValueError(f'BOQ import is limited to {_BOQ_IMPORT_MAX_ROWS} data rows')
    headers = [str(value or '').strip() for value in cleaned[0]]
    if not any(headers):
        raise ValueError('The selected sheet has no header row')
    records = []
    for source_row, values in enumerate(cleaned[1:], 2):
        records.append((source_row, {
            header: values[index] if index < len(values) else None
            for index, header in enumerate(headers) if header
        }))
    return headers, records


def _suggest_boq_mapping(headers):
    aliases = {
        'num': ('item number', 'item no', 'item', 'no', '#', 'رقم الفقرة', 'رقم البند'),
        'desc': ('description', 'work item description', 'item description', 'الوصف', 'وصف الفقرة'),
        'unit': ('unit', 'الوحدة'),
        'qty': ('quantity', 'qty', 'الكمية'),
        'price': ('unit price', 'rate', 'price', 'سعر الوحدة', 'السعر'),
        'total': ('total', 'amount', 'المجموع', 'الاجمالي', 'الإجمالي'),
    }
    normalized = {str(header).strip().casefold(): header for header in headers}
    result = {}
    for field, candidates in aliases.items():
        for candidate in candidates:
            if candidate.casefold() in normalized:
                result[field] = normalized[candidate.casefold()]
                break
    return result


def _validate_boq_rows(records, mapping):
    valid, rejected = [], []
    for source_row, raw in records:
        description = str(raw.get(mapping.get('desc')) or '').strip()
        unit = str(raw.get(mapping.get('unit')) or '').strip()
        item_number = str(raw.get(mapping.get('num')) or '').strip()
        quantity = _parse_localized_number(raw.get(mapping.get('qty')))
        price = _parse_localized_number(raw.get(mapping.get('price')))
        source_total = _parse_localized_number(raw.get(mapping.get('total'))) if mapping.get('total') else None
        errors = []
        if not description:
            errors.append('description')
        if quantity is None or quantity < 0:
            errors.append('quantity')
        if price is None or price < 0:
            errors.append('unit_price')
        if errors:
            rejected.append({'source_row': source_row, 'errors': errors})
            continue
        valid.append({
            'source_row': source_row,
            'num': item_number,
            'desc': description,
            'unit': unit,
            'qty': quantity,
            'price': price,
            'source_total': source_total,
            'calculated_total': quantity * price,
        })
    return valid, rejected


@financial_bp.route('/boq-import/preview', methods=['POST'])
@jwt_required()
def preview_boq_import(project_id):
    uid = int(get_jwt_identity())
    proj, mem = _requires_role(project_id, uid, 'member')
    if proj is None:
        return jsonify({'error': 'Project not found'}), 404
    if mem is None:
        return jsonify({'error': 'Financial edit permission is required'}), 403
    uploaded = request.files.get('file')
    if uploaded is None or not uploaded.filename:
        return jsonify({'error': 'Select a .xlsx or .csv file'}), 400
    extension = Path(uploaded.filename).suffix.lower()
    if extension not in {'.csv', '.xlsx'}:
        return jsonify({'error': 'Only .xlsx and .csv files are supported'}), 400
    content = uploaded.stream.read(_BOQ_IMPORT_MAX_BYTES + 1)
    if len(content) > _BOQ_IMPORT_MAX_BYTES:
        return jsonify({'error': 'BOQ import files are limited to 5 MB'}), 413
    try:
        rows = _read_csv_rows(content) if extension == '.csv' else _read_xlsx_rows(content)
        headers, records = _clean_sheet(rows)
        suggested = _suggest_boq_mapping(headers)
        supplied = request.form.get('mapping')
        mapping = json.loads(supplied) if supplied else suggested
        if not isinstance(mapping, dict):
            raise ValueError('Column mapping must be an object')
        required = {'desc', 'qty', 'price'}
        if not required.issubset(mapping) or any(mapping[key] not in headers for key in required):
            raise ValueError('Map description, quantity, and unit price columns')
        valid, rejected = _validate_boq_rows(records, mapping)
    except (ValueError, TypeError, json.JSONDecodeError, zipfile.BadZipFile, KeyError) as error:
        return jsonify({'error': str(error)}), 400
    return jsonify({
        'filename': uploaded.filename,
        'headers': headers,
        'mapping': mapping,
        'suggested_mapping': suggested,
        'valid_rows': valid,
        'rejected_rows': rejected,
        'limits': {'max_bytes': _BOQ_IMPORT_MAX_BYTES, 'max_rows': _BOQ_IMPORT_MAX_ROWS},
    }), 200
